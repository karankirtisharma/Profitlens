"""Bounded, non-AI extraction. Unknown columns and original files are preserved."""

import csv
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader

MAX_ROWS = 2000


def extract(content, filename, adapter="generic"):
    rows, pages = [], []
    if content.startswith(b"%PDF-"):
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 100:
            raise ValueError("Encrypted or over-100-page PDF is unsupported")
        pages = [
            {"page": i + 1, "text": p.extract_text(extraction_mode="layout") or ""}
            for i, p in enumerate(reader.pages)
        ]
        if sum(len(p["text"]) for p in pages) > 1_000_000:
            raise ValueError("Extracted PDF text exceeds prototype limit")
        if adapter == "fhdc_card":
            for page in pages:
                if not all(
                    word in page["text"].lower()
                    for word in ("supplier", "vat", "amount")
                ):
                    raise ValueError("PDF does not match the reviewed FHDC card layout")
                for line in page["text"].splitlines():
                    if not re.search(r"\d{2}/\d{2}/\d{4}", line):
                        continue
                    cells = re.split(r"\s{2,}", line.strip())
                    if len(cells) != 7:
                        raise ValueError(
                            "Unexpected FHDC columns; no partial extraction accepted"
                        )
                    row = dict(
                        zip(
                            (
                                "service_area",
                                "description",
                                "transaction_date",
                                "net_amount",
                                "vat",
                                "gross",
                                "supplier",
                            ),
                            cells,
                            strict=True,
                        )
                    )

                    def number(value):
                        value = value.replace(",", "").replace("£", "")
                        return Decimal(
                            "-" + value[1:-1] if value.startswith("(") else value
                        )

                    if number(row["net_amount"]) + number(row["vat"]) != number(
                        row["gross"]
                    ):
                        raise ValueError(
                            "Source net + VAT differs from gross; review source"
                        )
                    row["source_page"] = str(page["page"])
                    rows.append(row)
        elif adapter != "generic":
            raise ValueError("Unknown PDF adapter")
        if not any(p["text"].strip() for p in pages):
            raise ValueError("PDF has no selectable text; OCR is not implemented")
        fmt = "pdf"
    elif filename.lower().endswith(".xlsx"):
        with ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 1000
                or sum(i.file_size for i in archive.infolist()) > 20_000_000
            ):
                raise ValueError("Expanded workbook exceeds safe limit")
        book = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(book.worksheets) != 1:
                raise ValueError(
                    "Import one worksheet per file; multiple sheets require explicit selection"
                )
            iterator = book.active.iter_rows()
            headers = [str(c.value or "") for c in next(iterator)]
            for cells in iterator:
                if any(c.data_type == "f" for c in cells):
                    raise ValueError(
                        "Formula cells require a values-only export; formulas are not executed"
                    )
                values = [
                    c.value.isoformat()
                    if isinstance(c.value, (date, datetime))
                    else ""
                    if c.value is None
                    else str(c.value)
                    for c in cells
                ]
                if any(values):
                    rows.append(dict(zip(headers, values, strict=True)))
                if len(rows) > MAX_ROWS:
                    raise ValueError("Select a source subset of at most 2,000 rows")
            check_headers(headers)
        finally:
            book.close()
        fmt = "xlsx"
    elif filename.lower().endswith(".json"):
        data = json.loads(content, parse_float=str, parse_int=str)
        if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
            raise ValueError("JSON must be an array of record objects")
        rows = [
            {
                k: ""
                if v is None
                else v
                if isinstance(v, str)
                else json.dumps(v, ensure_ascii=False)
                for k, v in r.items()
            }
            for r in data
        ]
        fmt = "json"
    else:
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        check_headers(reader.fieldnames or [])
        for row in reader:
            if None in row or any(v is None for v in row.values()):
                raise ValueError("CSV row has missing or excess cells")
            rows.append(row)
            if len(rows) > MAX_ROWS:
                raise ValueError("Select a source subset of at most 2,000 rows")
        fmt = "csv"
    if len(rows) > MAX_ROWS:
        raise ValueError("Select a source subset of at most 2,000 rows")
    headers = list(dict.fromkeys(k for row in rows for k in row))
    if headers:
        check_headers(headers)
    if not rows and fmt != "pdf":
        raise ValueError("No source rows")
    return fmt, headers, rows, pages


def check_headers(headers):
    if (
        not headers
        or len(headers) > 80
        or len(set(headers)) != len(headers)
        or any(not h.strip() or len(h) > 120 for h in headers)
    ):
        raise ValueError(
            "Require 1–80 unique nonempty column names, at most 120 characters each"
        )


def suggestions(headers):
    aliases = {
        "amount": {"amount", "net_amount", "total", "value"},
        "date": {"date", "transaction_date", "invoicedate", "invoice_date"},
        "counterparty": {"supplier", "vendor", "customer", "counterparty"},
        "description": {"description", "memo", "details"},
        "external_id": {"transaction_id", "record_id"},
        "customer_id": {"customer_id", "customerid"},
        "currency": {"currency", "currency_code"},
    }
    return {
        target: {
            "column": matches[0] if len(matches) == 1 else "",
            "confidence": "candidate"
            if len(matches) == 1
            else "ambiguous"
            if matches
            else "unmapped",
            "candidates": matches,
        }
        for target, names in aliases.items()
        for matches in [
            [h for h in headers if h.lower().strip().replace(" ", "_") in names]
        ]
    }
